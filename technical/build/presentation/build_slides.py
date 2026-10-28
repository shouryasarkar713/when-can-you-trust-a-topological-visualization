#!/usr/bin/env python3
"""Build the final deck (PPTX with speaker notes) with python-pptx from slides_content.py.

    python build_slides.py <figures_dir> <out.pptx> <notes_md_out>

All text (slide text AND speaker notes) lives in slides_content.py; the notes file and the notes
embedded in the PPTX are written from the same objects, so they cannot differ.
"""
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

sys.path.insert(0, str(Path(__file__).parent))
from slides_content import SLIDES  # noqa: E402

NAVY = RGBColor(0x14, 0x2B, 0x4D)
INK = RGBColor(0x1F, 0x24, 0x2E)
GREY = RGBColor(0x5B, 0x64, 0x72)
BLUE = RGBColor(0x00, 0x72, 0xB2)
ORANGE = RGBColor(0xD5, 0x5E, 0x00)
GREEN = RGBColor(0x00, 0x9E, 0x73)
LIGHT = RGBColor(0xF2, 0xF5, 0xFA)
W, H = Inches(13.333), Inches(7.5)


def add_text(slide, x, y, w, h, paras, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, bullet=False, line=1.05):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    first = True
    for p_ in paras:
        if isinstance(p_, str):
            p_ = {"t": p_}
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        para.alignment = align
        para.line_spacing = line
        para.space_after = Pt(p_.get("after", 6))
        txt = p_["t"]
        if bullet and not p_.get("nobullet"):
            txt = "•  " + txt
        # **bold** segments
        parts = txt.split("**")
        for i, seg in enumerate(parts):
            if not seg:
                continue
            r = para.add_run()
            r.text = seg
            r.font.name = "Calibri"
            r.font.size = Pt(p_.get("size", size))
            r.font.bold = bold or (i % 2 == 1) or p_.get("bold", False)
            r.font.color.rgb = p_.get("color", color)
    return tb


def add_rect(slide, x, y, w, h, fill, line=None):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    s.adjustments[0] = 0.06
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    s.shadow.inherit = False
    return s


def add_pic(slide, path, x, y, w, h):
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(w / iw, h / ih)
    pw, ph = int(iw * scale), int(ih * scale)
    slide.shapes.add_picture(str(path), x + (w - pw) // 2, y + (h - ph) // 2, pw, ph)


def build(fig_dir, out_pptx, notes_out):
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]
    notes_md = ["# Speaker notes — When Can You Trust a Topological Visualization?\n"]
    total_words = 0
    for n, s in enumerate(SLIDES, 1):
        sl = prs.slides.add_slide(blank)
        if s.get("kind") == "title":
            bg = add_rect(sl, 0, 0, W, H, NAVY)
            bg.adjustments[0] = 0
            add_text(sl, Inches(0.9), Inches(2.0), Inches(11.5), Inches(1.8), [s["title"]], size=40, color=RGBColor(255, 255, 255), bold=True)
            add_text(sl, Inches(0.9), Inches(3.9), Inches(11.5), Inches(1.4), [{"t": s["subtitle"], "after": 10}], size=20, color=RGBColor(0xC9, 0xD6, 0xEA))
            add_text(sl, Inches(0.9), Inches(6.2), Inches(11.5), Inches(0.8), [s["footer"]], size=14, color=RGBColor(0xAF, 0xBF, 0xD8))
        else:
            add_text(sl, Inches(0.6), Inches(0.35), Inches(12.1), Inches(0.9), [s["title"]], size=30, color=NAVY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
            bar = add_rect(sl, Inches(0.65), Inches(1.22), Inches(1.2), Inches(0.06), BLUE)
            layout = s["layout"]
            body = s.get("body", [])
            if layout == "text":
                add_text(sl, Inches(0.7), Inches(1.5), Inches(11.9), Inches(5.2), body, size=s.get("size", 22), bullet=True)
            elif layout == "image_right":
                add_text(sl, Inches(0.7), Inches(1.5), Inches(5.0), Inches(5.3), body, size=s.get("size", 19), bullet=True)
                add_pic(sl, Path(fig_dir) / s["image"], Inches(5.8), Inches(1.5), Inches(7.1), Inches(5.0))
                if s.get("caption"):
                    add_text(sl, Inches(5.8), Inches(6.5), Inches(7.1), Inches(0.5), [s["caption"]], size=11, color=GREY, align=PP_ALIGN.CENTER)
            elif layout == "image_full":
                bh = s.get("body_h", 0.9)
                add_text(sl, Inches(0.7), Inches(1.4), Inches(11.9), Inches(bh), body, size=s.get("size", 18), bullet=s.get("bullets", False))
                add_pic(sl, Path(fig_dir) / s["image"], Inches(0.9), Inches(1.4 + bh + 0.1), Inches(11.5), Inches(6.7 - (1.4 + bh + 0.1)))
                if s.get("caption"):
                    add_text(sl, Inches(0.9), Inches(6.75), Inches(11.5), Inches(0.5), [s["caption"]], size=11, color=GREY, align=PP_ALIGN.CENTER)
            elif layout == "boxes":
                k = len(s["boxes"])
                bw = (12.0 - 0.25 * (k - 1)) / k
                for i, (head, txt, col) in enumerate(s["boxes"]):
                    x = Inches(0.65 + i * (bw + 0.25))
                    add_rect(sl, x, Inches(1.7), Inches(bw), Inches(3.3), LIGHT, line=col)
                    add_rect(sl, x, Inches(1.7), Inches(bw), Inches(0.62), col)
                    add_text(sl, x, Inches(1.72), Inches(bw), Inches(0.6), [head], size=17, color=RGBColor(255, 255, 255), bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
                    add_text(sl, x + Inches(0.08), Inches(2.45), Inches(bw - 0.16), Inches(2.5), [txt], size=s.get("size", 15))
                add_text(sl, Inches(0.7), Inches(5.3), Inches(11.9), Inches(1.7), body, size=19, bullet=True)
            elif layout == "two_image":
                add_text(sl, Inches(0.7), Inches(1.4), Inches(11.9), Inches(0.9), body, size=s.get("size", 18))
                add_pic(sl, Path(fig_dir) / s["image"], Inches(0.6), Inches(2.3), Inches(6.1), Inches(4.2))
                add_pic(sl, Path(fig_dir) / s["image2"], Inches(6.7), Inches(2.3), Inches(6.1), Inches(4.2))
                if s.get("caption"):
                    add_text(sl, Inches(0.7), Inches(6.6), Inches(11.9), Inches(0.5), [s["caption"]], size=11, color=GREY, align=PP_ALIGN.CENTER)
            add_text(sl, Inches(12.0), Inches(7.05), Inches(1.1), Inches(0.35), [f"{n}/{len(SLIDES)}"], size=10, color=GREY, align=PP_ALIGN.RIGHT)
        notes = s["notes"].strip()
        sl.notes_slide.notes_text_frame.text = notes
        total_words += len(notes.split())
        notes_md.append(f"\n## Slide {n} — {s.get('title')}\n\n{notes}\n")
    notes_md.append(f"\n---\nNarration word count: {total_words}; at 130 words per minute: {total_words/130:.1f} minutes.\n")
    prs.core_properties.title = "When Can You Trust a Topological Visualization?"
    prs.core_properties.author = "Shourya Sarkar"
    prs.core_properties.last_modified_by = "Shourya Sarkar"
    prs.core_properties.comments = ""
    prs.save(out_pptx)
    Path(notes_out).write_text("".join(notes_md), encoding="utf8")
    print(f"slides={len(SLIDES)} words={total_words} minutes@130wpm={total_words/130:.1f}")


if __name__ == "__main__":
    build(sys.argv[1], sys.argv[2], sys.argv[3])
