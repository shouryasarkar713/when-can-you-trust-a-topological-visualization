#!/usr/bin/env python3
"""Portable evidence verifier.

    python verify_evidence.py [--root PACKAGE_ROOT] [--out out.json] [--write-register]

Run it from ANY working directory (paths are resolved against --root, default: the parent of this
script's directory). It does not contain the report's numbers: it DERIVES each claimed quantity
from the result files (see facts.py) and then checks that the derived string appears in the
Markdown report, the generated DOCX, the generated PDF and, where listed, the slides and Q&A.

Checks: derived numeric claims; stale/forbidden phrases; figures (missing / orphaned);
citations (cited-but-not-listed / listed-but-not-cited); links; DOCX/PDF vs Markdown headings;
slide text vs embedded notes vs speaker_notes.md and the narration word count; environment
records (versions in the report vs requirements pins vs freeze); required outputs; placeholder text.
Exit status 0 only if every check passes.
"""
import argparse
import json
import re
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import facts as F  # noqa: E402

FORBIDDEN = [  # stale or unsupported wording that must not reappear in any deliverable
    "r₁ = 1.24", "right by luck", "longest basis cycle 4 nodes", "pre-registered", "preregistered 180",
    "Yvinec, B.", "15(5)", "r₁ ≥ 4.38", "TODO", "TBD", "lorem ipsum", "FIXME", "[placeholder", "XXX",
    "first Mapper explorer ever", "novel algorithm", "no existing tool", "frequently confuse", "routinely confuse",
    "pilot study was conducted", "participants were recruited", "fixed before outputs were inspected",
    "(0 in 6 ", "0 in six", "the annual signal lives",
]
FORBIDDEN_OK_CONTEXT = ["pre-registered", "pre-registration", "preregistered"]  # allowed only in negations (see check)


def squash(s):
    return re.sub(r"[\s|*`_]+", "", s).replace(" ", "")


def read_docx_text(p):
    from docx import Document
    d = Document(p)
    out = [para.text for para in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            out.append(" | ".join(c.text for c in row.cells))
    return "\n".join(out)


def read_pdf_text(p):
    import pymupdf
    with pymupdf.open(p) as d:
        return "\n".join(pg.get_text() for pg in d), len(d)


def read_pptx(p):
    from pptx import Presentation
    prs = Presentation(p)
    slides, notes = [], []
    for s in prs.slides:
        t = []
        for sh in s.shapes:
            if sh.has_text_frame:
                t.append(sh.text_frame.text)
        slides.append("\n".join(t))
        notes.append(s.notes_slide.notes_text_frame.text if s.has_notes_slide else "")
    return slides, notes


def main():
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    ap = argparse.ArgumentParser()
    ap.add_argument("--root", default=str(HERE.parent))
    ap.add_argument("--out", default=None)
    ap.add_argument("--write-register", action="store_true")
    a = ap.parse_args()
    root = Path(a.root).resolve()
    res = {"root": str(root), "checks": []}

    def check(name, ok, detail=""):
        res["checks"].append({"name": name, "pass": bool(ok), "detail": str(detail)[:400]})

    SUB, TECH = root / "submission", root / "technical"
    VER = TECH / "verification"
    rep_md = SUB / "report" / "Final_Report.md"
    rep_docx = SUB / "report" / "Final_Report.docx"
    rep_pdf = SUB / "report" / "Final_Report.pdf"
    pptx = SUB / "presentation" / "Final_Presentation.pptx"
    notes_md = SUB / "presentation" / "Speaker_Notes.md"
    qa_md = TECH / "build" / "presentation" / "QA_Preparation.md"
    qa_docx = SUB / "presentation" / "QA_Preparation.docx"
    qa_pdf = SUB / "presentation" / "QA_Preparation.pdf"

    required = [rep_md, rep_docx, rep_pdf, pptx, notes_md, qa_md, qa_docx, qa_pdf,
                root / "README.md", VER / "claims_register.json",
                VER / "requirements_matrix.md", VER / "research_ledger.md",
                VER / "code_audit_findings.md", TECH / "scikit-tda-trust" / "requirements.txt",
                TECH / "scikit-tda-trust" / "scripts" / "run_all_tests.py"]
    for p in required:
        check(f"required output exists: {p.relative_to(root).as_posix()}", p.exists())
    if not rep_md.exists():
        print(json.dumps(res["checks"], indent=1)); return 1

    md = rep_md.read_text(encoding="utf8")
    texts = {"report": {"md": md}}
    if rep_docx.exists():
        texts["report"]["docx"] = read_docx_text(rep_docx)
    if rep_pdf.exists():
        t, npages = read_pdf_text(rep_pdf)
        short = re.search(r"^short:\s*(.+)$", md, re.M)
        drop = {short.group(1).strip()} if short else set()
        t = chr(10).join(l for l in t.splitlines() if l.strip() not in drop and not re.match(r"^Page \d+$", l.strip()))  # running header/footer
        texts["report"]["pdf"] = t
        res["report_pdf_pages"] = npages
    slides_t, notes_t = ([], [])
    if pptx.exists():
        slides_t, notes_t = read_pptx(pptx)
    texts["slides"] = {"pptx": "\n".join(slides_t), "notes": "\n".join(notes_t)}
    texts["qa"] = {"md": qa_md.read_text(encoding="utf8") if qa_md.exists() else "",
                   "docx": read_docx_text(qa_docx) if qa_docx.exists() else "",
                   "pdf": read_pdf_text(qa_pdf)[0] if qa_pdf.exists() else ""}

    # 1. derived numeric claims
    register = []
    for fid, cat, claim, fn, where, cmd, cav in F.FACTS:
        try:
            val = fn(str(root))
        except Exception as e:  # a failed derivation is a failed check, never skipped
            check(f"{fid} derivation: {claim}", False, repr(e))
            continue
        sq = squash(val)
        ok_all, miss = True, []
        for w in where:
            for fmt, t in texts[w].items():
                if w == "slides" and fmt == "notes":
                    continue   # slide-specific claims are checked against slide text; notes are checked separately
                if t and sq not in squash(t):
                    ok_all = False
                    miss.append(f"{w}/{fmt}")
        check(f"{fid}: '{val[:70]}' derived from results and present in {','.join(where)}", ok_all,
              ("missing in " + ",".join(miss)) if miss else "")
        register.append({"id": fid, "category": cat, "claim": claim, "derived_value": val,
                         "documents": where, "reproduce": cmd, "caveat": cav,
                         "status": "verified" if ok_all else "FAILED"})
    res["n_facts"] = len(register)

    # 2. forbidden / stale wording (allowed only where it is quoted as a withdrawn/corrected statement or negated)
    QUOTE_MARKERS = re.compile(r"earlier|withdrawn|corrected|wrong by|was changed|stale|dropped|too strong|A-0\d|A-1\d|"
                               r"\b(not|never|no|nor|neither|without|nothing)\b", re.I)
    for phrase in FORBIDDEN:
        hits = []
        for w, d in texts.items():
            for fmt, t in d.items():
                for m in re.finditer(re.escape(phrase), t, re.I):
                    window = t[max(0, m.start() - 160): m.end() + 120]
                    if QUOTE_MARKERS.search(window):
                        continue
                    hits.append(f"{w}/{fmt}: ..." + t[max(0, m.start() - 60): m.end() + 40].replace(chr(10), " ") + "...")
        check(f"no stale/unsupported wording outside corrected/negated context: '{phrase}'", not hits, hits[:2])

    # 2b. process / history language must not appear in the submission-facing documents
    HISTORY = [r"previous (version|report|agent|solution|build|iteration|deck)", r"earlier (version|report|draft|deck|extension)",
               r"\bold (report|version|conclusion|deck)", r"independent rebuild", r"\brebuild", r"supersed", r"this iteration",
               r"(was|were) corrected", r"we fixed", r"code audit", r"audit record", r"\bagent\b", r"\bprompt\b", r"\bTODO\b",
               r"[A-Z]:\\Users", r"\bE:\\", r"/tmp/", r"independent_rebuild", r"final_submission", r"\bchangelog\b"]
    for pat in HISTORY:
        hits = []
        for w, d in texts.items():
            for fmt, t in d.items():
                for m in re.finditer(pat, t, re.I):
                    hits.append(f"{w}/{fmt}: ..." + t[max(0, m.start() - 40): m.end() + 40].replace(chr(10), " ") + "...")
        check(f"no project-history/process language: /{pat}/", not hits, hits[:2])

    # 3. figures
    figs = sorted(set(re.findall(r"!\[[^\]]*\]\(([^)]+)\)", md)))
    for fpath in figs:
        check(f"figure exists: {fpath}", (SUB / "report" / fpath).exists())
    fdir = SUB / "report" / "figures"
    listed = {Path(x).name for x in figs}
    orphans = [p.name for p in fdir.glob("*.png") if p.name not in listed] if fdir.exists() else []
    check("no orphaned figure files in report/figures", not orphans, orphans)
    for fpath in figs:
        n_caps = len(re.findall(r"^\*Figure [0-9.]+\.", md, re.M))
    n_caps = len(re.findall(r"^\*Figure [0-9]+\.[0-9]+\.", md, re.M))
    check("every figure has a caption", len(figs) == n_caps, "%d figures, %d captions" % (len(figs), n_caps))

    # 4. citations
    refs_sec = md.split("## 13. References", 1)[1].split("## Appendix A", 1)[0] if "## 13. References" in md else ""
    ref_entries = [l[2:] for l in refs_sec.splitlines() if l.startswith("- ")]
    ref_keys = set()
    for e_ in ref_entries:
        m = re.match(r"([^,(]+?)[,.].*?\((\d{4})\)", e_)
        if m:
            ref_keys.add((m.group(1).strip().split()[-1], m.group(2)))
        else:
            m2 = re.match(r"(NOAA[^.]*)\.", e_)
    body = md.split("## 13. References", 1)[0]
    cited = set()
    for m in re.finditer(r"([A-Z][A-Za-zÀ-ÿ'’\-]+)(?: et al\.|(?:,| &) [A-Z][A-Za-zÀ-ÿ'’\-]+(?:(?:,| &| and) [A-Z][A-Za-zÀ-ÿ'’\-]+)*)?(?:,)? \(?(\d{4})\)?", body):
        cited.add((m.group(1), m.group(2)))
    known_authors = {k[0] for k in ref_keys}
    missing = sorted(c for c in cited if c[0] in known_authors and c not in ref_keys)
    check("every Author (Year) citation of a listed author matches a reference entry", not missing, missing)
    # reverse: each listed reference cited at least once in the body (except dataset entries)
    uncited = []
    for k in sorted(ref_keys):
        if k[0] in ("NOAA",):
            continue
        if not re.search(re.escape(k[0]) + r"[^\n]{0,70}\(?" + k[1], body):
            uncited.append(k)
    check("every listed reference is cited in the body", not uncited, uncited)
    check("reference list has no duplicate (author, year, title) entries", len(ref_entries) == len(set(ref_entries)))

    # 5. links
    for m in re.finditer(r"\]\((?!http)([^)#]+)\)", md):
        tgt = m.group(1)
        if tgt.startswith(("mailto:",)):
            continue
        check(f"relative link/figure resolves: {tgt}", (SUB / "report" / tgt).exists())
    for u in set(re.findall(r"https?://[^\s)>\]]+", md)):
        ok = re.match(r"^https?://[A-Za-z0-9.\-]+(/[^\s]*)?$", u.rstrip(".,;"))
        check(f"URL well-formed: {u[:80]}", ok)

    # 6. DOCX / PDF vs Markdown headings
    heads = [re.sub(r"[*`]", "", h).strip() for h in re.findall(r"^#{2,3} (.+)$", md, re.M)]
    for fmt in ("docx", "pdf"):
        t = texts["report"].get(fmt, "")
        miss = [h for h in heads if squash(h) not in squash(t)]
        check(f"all {len(heads)} Markdown headings present in {fmt}", not miss, miss[:3])
    if "pdf" in texts["report"]:
        check("PDF has at least 20 pages (sanity bound) and a rendered table of contents",
              res.get("report_pdf_pages", 0) >= 20 and "Contents" in texts["report"]["pdf"], res.get("report_pdf_pages"))

    # 6b. staleness of generated editions: every Markdown paragraph / list item (complete text)
    # must appear in the DOCX and in the PDF
    items, cur, in_code = [], [], False

    def flush():
        if cur:
            items.append(" ".join(cur))
            cur.clear()
    in_comment = False
    for ln in md.splitlines():
        if ln.lstrip().startswith("<!--") and "-->" not in ln:
            in_comment = True
            continue
        if in_comment:
            if "-->" in ln:
                in_comment = False
            continue
        if ln.strip().startswith("```"):
            in_code = not in_code
            flush()
            continue
        if in_code or not ln.strip() or ln.lstrip().startswith(("#", "|", "!", "<!--", "-->")):
            flush()
            continue
        marker = re.match(r"^\s*(?:[-*]|\d+\.)\s+(.*)", ln)
        if marker:
            flush()
            cur.append(marker.group(1))
        else:
            cur.append(ln.strip())
    flush()
    paras = [re.sub(r"<(https?://[^>\s]+)>", r"\1", re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", p)) for p in items]
    for fmt in ("docx", "pdf"):
        t = squash(texts["report"].get(fmt, ""))
        stale = [p[:60] for p in paras if squash(p) not in t]
        check(f"complete text of all {len(paras)} Markdown paragraphs/items present in {fmt} (edition not stale)",
              not stale, stale[:3])

    # 7. slides / notes
    if slides_t:
        n_md = re.split(r"^## Slide \d+ — .*$", notes_md.read_text(encoding="utf8"), flags=re.M)[1:]
        n_md = [re.sub(r"\n---\nNarration.*$", "", x, flags=re.S).strip() for x in n_md]
        check("slide count == notes sections", len(slides_t) == len(n_md), f"{len(slides_t)} vs {len(n_md)}")
        same = [i + 1 for i, (x, y) in enumerate(zip(notes_t, n_md)) if x.strip() != y.strip()]
        check("embedded PPTX notes identical to speaker_notes.md", not same, same)
        check("every slide has speaker notes of 60-200 words",
              all(60 <= len(x.split()) <= 200 for x in notes_t), [len(x.split()) for x in notes_t])
        words = sum(len(x.split()) for x in notes_t)
        m = re.search(r"Narration word count: (\d+); at 130 words per minute: ([0-9.]+) minutes", notes_md.read_text(encoding="utf8"))
        check("documented narration duration equals actual word count / 130 wpm", m and int(m.group(1)) == words and abs(float(m.group(2)) - words / 130) < 0.06, (m.groups() if m else None))
        check("talk length 9-12 minutes", 9 <= words / 130 <= 12, f"{words/130:.1f} min")
        for fid, cat, claim, fn, where, cmd, cav in F.FACTS:
            pass

    # 8. environment records
    req = (TECH / "scikit-tda-trust" / "requirements.txt").read_text(encoding="utf8")
    pins = dict(re.findall(r"^([A-Za-z0-9_\-]+)==([0-9.]+)", req, re.M))
    for pkg, ver in [("ripser", "0.6.15"), ("persim", "0.4.1"), ("kmapper", "2.1.0"), ("numpy", "2.5.4"), ("scipy", "1.18.1"),
                     ("scikit-learn", "1.9.1"), ("matplotlib", "3.11.2"), ("networkx", "3.7")]:
        check(f"requirements pin {pkg}=={ver}", pins.get(pkg) == ver, pins.get(pkg))
        short = {"scikit-learn": "scikit-learn", "kmapper": "kmapper"}.get(pkg, pkg)
        check(f"report states {pkg} {ver}", f"{short} {ver}" in md or f"**{short} {ver}**" in md or f"{short} {ver}" in md.replace("**", ""), "")
    fz = VER / "env" / "pip_freeze_clean_environment.txt"
    if fz.exists():
        frz = dict(re.findall(r"^([A-Za-z0-9_.\-]+)==([0-9.]+)", fz.read_text(encoding="utf8"), re.M))
        for pkg, ver in pins.items():
            check(f"environment freeze matches pin {pkg}", frz.get(pkg, frz.get(pkg.replace('-', '_'), frz.get(pkg.lower()))) == ver, frz.get(pkg))
    else:
        check("environment freeze file present", False, str(fz))

    # 8b. repository link and version-control disclosure
    REPO_URL = "https://github.com/shouryasarkar713/when-can-you-trust-a-topological-visualization"
    for fmt in ("md", "docx", "pdf"):
        t = texts["report"].get(fmt, "")
        check(f"report {fmt}: contains the repository URL", squash(REPO_URL) in squash(t))
    if rep_pdf.exists():
        import pymupdf
        with pymupdf.open(rep_pdf) as pdf_doc:
            uris = {l.get("uri") for pg in pdf_doc for l in pg.get_links() if l.get("uri")}
        check("report pdf: the repository URL is a clickable link annotation", any(u and u.rstrip("/") == REPO_URL for u in uris),
              sorted(u for u in uris if "github" in u))
    if rep_docx.exists():
        import zipfile as _zf
        with _zf.ZipFile(rep_docx) as zf:
            rels = zf.read("word/_rels/document.xml.rels").decode("utf8", "replace")
        check("report docx: the repository URL is a hyperlink relationship", REPO_URL in rels)
    DISCLOSURE = "Contemporaneous version-control history was not maintained during the project."
    RETRO = "its commit history should not be interpreted as evidence of the original design or analysis chronology"
    for fmt in ("md", "docx", "pdf"):
        t = squash(texts["report"].get(fmt, ""))
        check(f"report {fmt}: version-control disclosure present", squash(DISCLOSURE) in t and squash(RETRO) in t)
    check("README states that the repository was organised retrospectively",
          "retrospectively" in (root / "README.md").read_text(encoding="utf8"))

    # 9. placeholders / emptiness in generated documents
    for w in ("docx", "pdf"):
        t = texts["report"].get(w, "")
        check(f"report {w}: no unresolved field-code or placeholder text", not re.search(r"Error! |\[Author\]|<<|>>|\{\{|\}\}", t))

    res["summary"] = {"n_checks": len(res["checks"]), "n_pass": sum(c["pass"] for c in res["checks"]),
                      "failed": [c for c in res["checks"] if not c["pass"]]}
    if a.write_register:
        VER.mkdir(exist_ok=True)
        manual = json.loads((HERE / "claims_manual.json").read_text(encoding="utf8")) if (HERE / "claims_manual.json").exists() else []
        note = ("Each claim has a stable id, category (measured / reproduced / external / interpretation / future), evidence "
                "path(s) relative to the package root, the command that reproduces it and its verification status. "
                "F* entries are re-derived from result files by verify_evidence.py on every run; M* entries are curated.")
        (VER / "claims_register.json").write_text(
            json.dumps({"note": note, "claims": manual + register}, indent=1, ensure_ascii=False), encoding="utf8")
    # Routine runs never write into the package: the detailed result is written only when --out is given
    # (the test runner passes a path in its own output directory).
    if a.out:
        Path(a.out).write_text(json.dumps(res, indent=1, ensure_ascii=False), encoding="utf8")
    print(f"{res['summary']['n_pass']}/{res['summary']['n_checks']} checks passed; {res['n_facts']} derived facts")
    for c in res["summary"]["failed"]:
        print("  FAIL:", c["name"], "|", c["detail"][:160])
    return 0 if not res["summary"]["failed"] else 1


if __name__ == "__main__":
    sys.exit(main())
